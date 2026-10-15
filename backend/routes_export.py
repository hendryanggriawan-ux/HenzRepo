from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core import OWNER, TZ, db, local_range, today_local
from models import CashClosing
from routes_master import get_settings
from routes_reports import mechanics_report, purchases_report, sales_report, stock_report
from routes_trx import list_moves

router = APIRouter(prefix="/export")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
METHOD_LABEL = {"tunai": "Tunai", "qris": "QRIS", "debit": "Debit", "kredit": "Kartu Kredit", "transfer": "Transfer"}
TYPE_LABEL = {"part": "Penjualan Part", "service": "Servis"}
HEAD_FILL = PatternFill("solid", fgColor="F97316")


def _dt(value: str):
    return datetime.fromisoformat(value).astimezone(TZ).replace(tzinfo=None) if value else None


def _period(start: str, end: str):
    s = start or today_local().isoformat()
    return s, end or s


def _new_book() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _add_sheet(wb, name, title, subtitle, headers, rows, money=()):
    ws = wb.create_sheet(name[:31])
    ws.append([title])
    ws.append([subtitle])
    ws.append([])
    ws.append(headers)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"].font = Font(italic=True, color="555555")
    head = ws.max_row
    for cell in ws[head]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(list(r))
    money_cols = {i + 1 for i, h in enumerate(headers) if h in money}
    widths = [len(str(h)) for h in headers]
    for row in ws.iter_rows(min_row=head + 1):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"
            elif cell.column in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
            if cell.column <= len(widths):
                size = 16 if isinstance(cell.value, datetime) else len(str(cell.value if cell.value is not None else ""))
                widths[cell.column - 1] = max(widths[cell.column - 1], size)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), 50)
    ws.freeze_panes = ws.cell(row=head + 1, column=1)


def _respond(wb: Workbook, filename: str) -> StreamingResponse:
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{filename}"'})


async def _subtitle(text: str) -> str:
    return f"{(await get_settings()).store_name} · {text} · diekspor {datetime.now(TZ).strftime('%d/%m/%Y %H:%M')}"


@router.get("/sales")
async def export_sales(start: str = "", end: str = "", type: str = "", _=Depends(OWNER)):
    s, e = _period(start, end)
    data = await sales_report(s, e, type, None)
    sm, rows = data["summary"], data["rows"]
    sub = await _subtitle(f"Periode {s} s/d {e}" + (f" · {TYPE_LABEL[type]}" if type in TYPE_LABEL else ""))
    wb = _new_book()
    summary = [
        ("Jumlah Transaksi", sm["count"]), ("Transaksi Penjualan Part", sm["part_count"]), ("Transaksi Servis", sm["service_count"]),
        ("Penjualan Part (Rp)", sm["parts_total"]), ("Pendapatan Jasa (Rp)", sm["services_total"]), ("Diskon (Rp)", sm["discount"]),
        ("Total Penjualan (Rp)", sm["total"]), ("HPP Part (Rp)", sm["cost_total"]), ("Laba Kotor (Rp)", sm["gross_profit"]),
    ] + [(f"Metode {label} (Rp)", sm["by_method"].get(m, 0)) for m, label in METHOD_LABEL.items()]
    _add_sheet(wb, "Ringkasan", "Laporan Penjualan", sub, ["Keterangan", "Nilai"], summary, {"Nilai"})
    money = {"Part", "Jasa", "Subtotal", "Diskon", "Total", "Bayar", "Kembali", "HPP"}
    _add_sheet(wb, "Transaksi", "Daftar Transaksi", sub,
               ["No. Transaksi", "Waktu", "Jenis", "No. WO", "No. Polisi", "Pelanggan", "No. HP", "Mekanik", "Kasir",
                "Metode", "No. Ref", "Part", "Jasa", "Subtotal", "Diskon", "Total", "Bayar", "Kembali", "HPP"],
               [(x.number, _dt(x.created_at), TYPE_LABEL[x.type], x.work_order_number, x.plate, x.customer_name,
                 x.customer_phone, x.mechanic_name, x.cashier, METHOD_LABEL.get(x.payment_method, x.payment_method),
                 x.payment_ref, x.parts_total, x.services_total, x.subtotal, x.discount, x.total, x.paid_amount,
                 x.change, x.cost_total) for x in rows], money)
    detail = []
    for x in rows:
        detail += [(x.number, _dt(x.created_at), "Part", i.code, i.name, i.qty, i.price, i.subtotal, x.mechanic_name) for i in x.items]
        detail += [(x.number, _dt(x.created_at), "Jasa", "", v.name, 1, v.price, v.price, x.mechanic_name) for v in x.services]
    _add_sheet(wb, "Detail Item", "Detail Part & Jasa Terjual", sub,
               ["No. Transaksi", "Waktu", "Jenis", "Kode", "Nama", "Qty", "Harga", "Subtotal", "Mekanik"], detail, {"Harga", "Subtotal"})
    return _respond(wb, f"laporan-penjualan_{s}_{e}.xlsx")


@router.get("/purchases")
async def export_purchases(start: str = "", end: str = "", _=Depends(OWNER)):
    s, e = _period(start, end)
    data = await purchases_report(s, e, None)
    sm, rows = data["summary"], data["rows"]
    sub = await _subtitle(f"Periode {s} s/d {e}")
    wb = _new_book()
    summary = [("Jumlah Faktur", sm["count"]), ("Total Pembelian (Rp)", sm["total"]), ("Qty Barang Masuk", sm["items_qty"])]
    summary += [(f"Supplier: {name} (Rp)", total) for name, total in sm["by_supplier"].items()]
    _add_sheet(wb, "Ringkasan", "Laporan Pembelian", sub, ["Keterangan", "Nilai"], summary, {"Nilai"})
    _add_sheet(wb, "Pembelian", "Daftar Pembelian", sub,
               ["No. Pembelian", "Waktu Input", "Tgl Faktur", "Supplier", "No. Faktur", "Jumlah Item", "Total", "Catatan", "Oleh"],
               [(p.number, _dt(p.created_at), p.date, p.supplier_name, p.invoice_no, len(p.items), p.total, p.note, p.created_by) for p in rows],
               {"Total"})
    detail = [(p.number, p.date, p.supplier_name, i.code, i.name, i.qty, i.price, i.subtotal) for p in rows for i in p.items]
    _add_sheet(wb, "Detail Item", "Detail Barang Masuk", sub,
               ["No. Pembelian", "Tgl Faktur", "Supplier", "Kode", "Nama", "Qty", "Harga Beli", "Subtotal"], detail, {"Harga Beli", "Subtotal"})
    return _respond(wb, f"laporan-pembelian_{s}_{e}.xlsx")


@router.get("/stock")
async def export_stock(mode: str = "all", _=Depends(OWNER)):
    data = await stock_report(mode, None)
    low, high = data["low_threshold"], data["high_threshold"]
    titles = {"low": "Laporan Stok Menipis", "high": "Laporan Stok Masih Banyak"}

    def status(i):
        if i.stock <= (i.min_stock if i.min_stock > 0 else low):
            return "Menipis"
        return "Banyak" if i.stock >= high else "Normal"

    sub = await _subtitle(f"Batas menipis {low:g} · batas banyak {high:g}")
    wb = _new_book()
    _add_sheet(wb, "Stok", titles.get(data["mode"], "Laporan Stok Barang"), sub,
               ["Kode", "Nama Barang", "Kategori", "Satuan", "Stok", "Stok Min", "Harga Beli", "Harga Jual",
                "Nilai Stok (Beli)", "Nilai Stok (Jual)", "Status"],
               [(i.code, i.name, i.category, i.unit, i.stock, i.min_stock, i.buy_price, i.sell_price,
                 round(i.stock * i.buy_price, 2), round(i.stock * i.sell_price, 2), status(i)) for i in data["rows"]],
               {"Harga Beli", "Harga Jual", "Nilai Stok (Beli)", "Nilai Stok (Jual)"})
    return _respond(wb, f"laporan-stok-{data['mode']}_{today_local().isoformat()}.xlsx")


@router.get("/mechanics")
async def export_mechanics(start: str = "", end: str = "", _=Depends(OWNER)):
    s, e = _period(start, end)
    data = await mechanics_report(s, e, None)
    sub = await _subtitle(f"Periode {s} s/d {e}")
    wb = _new_book()
    money = {"Pendapatan Jasa", "Penjualan Part", "Total", "Jasa (Rp)", "Part (Rp)"}
    _add_sheet(wb, "Ringkasan Mekanik", "Laporan Jasa Mekanik", sub,
               ["Mekanik", "Servis Selesai", "Servis Aktif", "Pendapatan Jasa", "Penjualan Part", "Total"],
               [(m["mechanic_name"], m["jobs"], m["active_jobs"], m["services_total"], m["parts_total"], m["total"]) for m in data["summary"]],
               money)
    _add_sheet(wb, "Pekerjaan", "Daftar Pekerjaan Mekanik", sub,
               ["Waktu", "No. Transaksi", "No. WO", "No. Polisi", "Kendaraan", "Pelanggan", "Mekanik", "Jasa", "Jasa (Rp)", "Part (Rp)", "Total"],
               [(_dt(j["created_at"]), j["number"], j["work_order_number"], j["plate"], j["vehicle_desc"], j["customer_name"],
                 j["mechanic_name"], ", ".join(j["services"]), j["services_total"], j["parts_total"], j["total"]) for j in data["jobs"]],
               money)
    return _respond(wb, f"laporan-jasa-mekanik_{s}_{e}.xlsx")


@router.get("/moves")
async def export_moves(start: str = "", end: str = "", q: str = "", _=Depends(OWNER)):
    s, e = _period(start, end)
    rows = await list_moves(s, e, q, None)
    wb = _new_book()
    _add_sheet(wb, "Mutasi Stok", "Laporan Mutasi Stok", await _subtitle(f"Periode {s} s/d {e}"),
               ["Waktu", "Kode", "Barang", "Jenis", "Qty", "Stok Akhir", "Referensi", "Keterangan", "User"],
               [(_dt(m.created_at), m.code, m.name, "Masuk" if m.type == "in" else "Keluar", m.qty, m.stock_after, m.ref, m.note, m.user) for m in rows])
    return _respond(wb, f"mutasi-stok_{s}_{e}.xlsx")


@router.get("/closings")
async def export_closings(start: str = "", end: str = "", _=Depends(OWNER)):
    s, e = _period(start, end)
    docs = await db.cash_closings.find({"created_at": local_range(s, e)}).sort("created_at", -1).to_list(5000)
    rows = [CashClosing.from_mongo(d) for d in docs]
    money = {"Total Penjualan", *METHOD_LABEL.values(), "Modal Awal", "Kas Seharusnya", "Uang Fisik", "Selisih"}
    wb = _new_book()
    _add_sheet(wb, "Tutup Kasir", "Laporan Tutup Kasir", await _subtitle(f"Periode {s} s/d {e}"),
               ["No", "Waktu Tutup", "Kasir", "Periode Mulai", "Transaksi", "Total Penjualan", *METHOD_LABEL.values(),
                "Modal Awal", "Kas Seharusnya", "Uang Fisik", "Selisih", "Catatan"],
               [(c.number, _dt(c.period_end), c.user_name, _dt(c.period_start), c.count, c.total,
                 *[c.by_method.get(m, 0) for m in METHOD_LABEL], c.opening_cash, c.expected_cash, c.counted_cash,
                 c.difference, c.note) for c in rows],
               money)
    return _respond(wb, f"tutup-kasir_{s}_{e}.xlsx")
