import re
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook

from core import OWNER, db, now_iso
from models import Item
from routes_export import _add_sheet, _new_book, _respond
from seed import gen_item_code
from stock import log_move

router = APIRouter(prefix="/items")
MAX_BYTES = 5_000_000
MAX_ROWS = 5000
HEADERS = ["Kode / Barcode", "Nama Barang *", "Kategori", "Satuan", "Harga Beli", "Harga Jual *", "Stok Awal", "Stok Minimum"]
LABEL = {"buy_price": "Harga Beli", "sell_price": "Harga Jual", "stock": "Stok Awal", "min_stock": "Stok Minimum"}
GUIDE = [
    ("Kode / Barcode", "Kosongkan agar dibuat otomatis, atau isi barcode pabrik. Jika kode sudah terdaftar, data barang tersebut DIPERBARUI."),
    ("Nama Barang *", "Wajib diisi."),
    ("Kategori", "Opsional, contoh: Oli, Rem, Kelistrikan."),
    ("Satuan", "Opsional, default pcs."),
    ("Harga Beli", "Angka tanpa Rp/titik, contoh 45000."),
    ("Harga Jual *", "Wajib diisi, angka tanpa Rp/titik, contoh 55000."),
    ("Stok Awal", "Hanya untuk barang BARU. Stok barang yang sudah ada tidak diubah (gunakan Pembelian / Mutasi Stok)."),
    ("Stok Minimum", "Batas peringatan stok menipis."),
    ("Catatan", "Isi data mulai baris 5 di sheet Barang. Hapus 2 baris contoh sebelum import. Kolom Kode diformat teks agar angka 0 di depan tidak hilang."),
]


def _field(header):
    h = " ".join(re.sub(r"[^a-z ]", " ", str(header or "").lower()).split())
    if not h:
        return None
    if h.startswith("kode") or "barcode" in h:
        return "code"
    if h.startswith("nama"):
        return "name"
    if h.startswith("kategori"):
        return "category"
    if h.startswith("satuan"):
        return "unit"
    if "beli" in h:
        return "buy_price"
    if "jual" in h:
        return "sell_price"
    if "min" in h:
        return "min_stock"
    if h.startswith("stok"):
        return "stock"
    return None


def _num(value) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().lower().replace("rp", "").replace(" ", "")
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", s):
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    return float(s)


def _code(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip() if value is not None else ""


def _clean(rec: dict) -> dict:
    name = str(rec.get("name") or "").strip()
    if not name:
        raise ValueError("Nama barang wajib diisi")
    if rec.get("sell_price") in (None, ""):
        raise ValueError("Harga jual wajib diisi")
    vals = {}
    for key, label in LABEL.items():
        try:
            vals[key] = _num(rec.get(key))
        except ValueError:
            raise ValueError(f"{label} bukan angka")
        if vals[key] < 0:
            raise ValueError(f"{label} tidak boleh negatif")
    return {"code": _code(rec.get("code")), "name": name, "category": str(rec.get("category") or "").strip(),
            "unit": str(rec.get("unit") or "").strip() or "pcs", **vals}


async def _parse(file: UploadFile):
    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(400, "Ukuran file maksimal 5MB")
    try:
        wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(400, "File tidak valid. Gunakan file Excel .xlsx")
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    head = next((i for i, r in enumerate(rows[:10]) if any(_field(c) == "name" for c in r)), None)
    if head is None:
        raise HTTPException(400, "Kolom 'Nama Barang' tidak ditemukan. Gunakan template import.")
    mapping = {j: _field(c) for j, c in enumerate(rows[head]) if _field(c)}
    out = []
    for no, r in enumerate(rows[head + 1:], start=head + 2):
        if all(c in (None, "") for c in r):
            continue
        out.append((no, {f: (r[j] if j < len(r) else None) for j, f in mapping.items()}))
    if len(out) > MAX_ROWS:
        raise HTTPException(400, f"Maksimal {MAX_ROWS} baris per import")
    return out


@router.get("/import-template")
async def import_template(_=Depends(OWNER)):
    wb = _new_book()
    _add_sheet(wb, "Barang", "Template Import Barang", "Isi mulai baris 5. Kolom bertanda * wajib. Hapus baris contoh sebelum import.",
               HEADERS, [("", "Contoh: Oli Mesin 1L", "Oli", "botol", 45000, 55000, 10, 5),
                         ("8991234567890", "Contoh: Busi Iridium", "Kelistrikan", "pcs", 60000, 85000, 5, 2)],
               {"Harga Beli", "Harga Jual *"})
    ws = wb["Barang"]
    for r in range(5, 2005):
        ws.cell(row=r, column=1).number_format = "@"
    _add_sheet(wb, "Petunjuk", "Petunjuk Pengisian", "Baca sebelum mengisi template", ["Kolom", "Keterangan"], GUIDE)
    return _respond(wb, "template-import-barang.xlsx")


async def _save(item: dict, action: str, user: dict) -> str:
    if action == "update":
        fields = {k: item[k] for k in ("name", "category", "unit", "buy_price", "sell_price", "min_stock")}
        await db.items.update_one({"code": item["code"]}, {"$set": fields})
        return item["code"]
    doc = Item(**{**item, "code": item["code"] or await gen_item_code()}, created_at=now_iso())
    res = await db.items.insert_one(doc.to_mongo())
    if doc.stock > 0:
        await log_move({**doc.to_mongo(), "_id": res.inserted_id}, "in", doc.stock, "IMPORT", "Stok awal (import Excel)", user["name"])
    return doc.code


@router.post("/import")
async def import_items(dry_run: bool = True, file: UploadFile = File(...), user: dict = Depends(OWNER)):
    parsed = await _parse(file)
    existing = {d["code"] async for d in db.items.find({}, {"code": 1})}
    seen, results = set(), []
    counts = {"baru": 0, "update": 0, "error": 0}
    for no, rec in parsed:
        try:
            item = _clean(rec)
            if item["code"] and item["code"] in seen:
                raise ValueError("Kode duplikat di dalam file")
        except ValueError as e:
            counts["error"] += 1
            results.append({"row": no, "code": _code(rec.get("code")), "name": str(rec.get("name") or ""), "action": "error", "message": str(e)})
            continue
        if item["code"]:
            seen.add(item["code"])
        action = "update" if item["code"] in existing else "baru"
        message = "Stok tidak diubah (barang sudah ada)" if action == "update" and item["stock"] else ""
        code = item["code"] if dry_run else await _save(item, action, user)
        counts[action] += 1
        results.append({"row": no, "code": code or "(otomatis)", "name": item["name"], "action": action, "message": message})
    return {"dry_run": dry_run, "total": len(parsed), "created": counts["baru"], "updated": counts["update"],
            "errors": counts["error"], "rows": results}
