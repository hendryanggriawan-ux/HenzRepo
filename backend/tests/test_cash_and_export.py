"""Tests for daily cash closing (Tutup Kasir) and Excel export endpoints."""
import io
import uuid
import requests
from datetime import date

from openpyxl import load_workbook

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# --------- helpers ----------
def _get_item(client, base_url, code):
    return client.get(f"{base_url}/api/items/code/{code}").json()


def _make_part_sale(kasir_client, base_url, code, qty=1, method="tunai"):
    item = _get_item(kasir_client, base_url, code)
    payload = {
        "items": [{"item_id": item["id"], "qty": qty}],
        "customer_name": "TEST_Cash",
        "payment": {"method": method, "ref": "TEST" if method != "tunai" else "",
                    "paid_amount": item["sell_price"] * qty if method == "tunai" else 0,
                    "discount": 0},
    }
    r = kasir_client.post(f"{base_url}/api/sales", json=payload)
    assert r.status_code == 200, r.text
    return r.json()


# ================================================================
# Cash closing (preview, close, closings list/get, isolation)
# ================================================================
def test_cash_preview_structure(kasir_client, base_url):
    r = kasir_client.get(f"{base_url}/api/cash/preview")
    assert r.status_code == 200
    data = r.json()
    for k in ("period_start", "period_end", "count", "part_count",
              "service_count", "total", "by_method", "cash_sales", "last_closing"):
        assert k in data, f"missing key {k}"
    for m in ("tunai", "qris", "debit", "kredit", "transfer"):
        assert m in data["by_method"]


def test_cash_preview_reflects_new_sale_and_close_flow(kasir_client, base_url):
    """Kasir makes a sale, preview shows it, closes, next preview is reset to 0."""
    pre = kasir_client.get(f"{base_url}/api/cash/preview").json()
    base_count = pre["count"]
    base_tunai = pre["by_method"]["tunai"]

    sale = _make_part_sale(kasir_client, base_url, "200000000001", qty=1, method="tunai")
    sale_total = sale["total"]

    after = kasir_client.get(f"{base_url}/api/cash/preview").json()
    assert after["count"] == base_count + 1
    assert round(after["by_method"]["tunai"] - base_tunai, 2) == sale_total
    assert after["cash_sales"] == after["by_method"]["tunai"]

    # Close
    opening = 100000.0
    expected = round(opening + after["cash_sales"], 2)
    counted = expected  # exact
    r = kasir_client.post(f"{base_url}/api/cash/close", json={
        "opening_cash": opening, "counted_cash": counted, "note": "TEST close"
    })
    assert r.status_code == 200, r.text
    cl = r.json()
    assert cl["number"].startswith("TK-")
    assert cl["expected_cash"] == expected
    assert cl["counted_cash"] == counted
    assert cl["difference"] == 0.0
    assert cl["user_name"]  # linked to user
    closing_id = cl["id"]

    # After closing, preview should reset for THIS user
    reset = kasir_client.get(f"{base_url}/api/cash/preview").json()
    assert reset["count"] == 0
    assert reset["total"] == 0
    assert reset["last_closing"]["id"] == closing_id

    # list contains it
    lst = kasir_client.get(f"{base_url}/api/cash/closings").json()
    assert any(c["id"] == closing_id for c in lst)

    # GET single
    r2 = kasir_client.get(f"{base_url}/api/cash/closings/{closing_id}")
    assert r2.status_code == 200
    assert r2.json()["id"] == closing_id


def test_cash_isolation_between_kasir_and_owner(kasir_client, owner_client, base_url):
    """Owner's preview should not include kasir's sales."""
    owner_pre = owner_client.get(f"{base_url}/api/cash/preview").json()
    kasir_pre = kasir_client.get(f"{base_url}/api/cash/preview").json()
    _make_part_sale(kasir_client, base_url, "200000000002", qty=1, method="qris")
    owner_post = owner_client.get(f"{base_url}/api/cash/preview").json()
    kasir_post = kasir_client.get(f"{base_url}/api/cash/preview").json()
    assert owner_post["count"] == owner_pre["count"], "kasir sale leaked into owner preview"
    assert kasir_post["count"] == kasir_pre["count"] + 1


def test_kasir_cannot_get_other_users_closing(kasir_client, owner_client, base_url):
    # Owner creates a closing of its own
    r = owner_client.post(f"{base_url}/api/cash/close", json={
        "opening_cash": 50000, "counted_cash": 50000, "note": "TEST owner close"
    })
    assert r.status_code == 200
    cid = r.json()["id"]
    # Kasir trying to fetch owner's closing → 404
    r2 = kasir_client.get(f"{base_url}/api/cash/closings/{cid}")
    assert r2.status_code == 404

    # Kasir list does not include owner closing
    lst = kasir_client.get(f"{base_url}/api/cash/closings").json()
    assert not any(c["id"] == cid for c in lst)

    # Owner list DOES include it
    olist = owner_client.get(f"{base_url}/api/cash/closings").json()
    assert any(c["id"] == cid for c in olist)


def test_sale_has_cashier_id(kasir_client, base_url):
    sale = _make_part_sale(kasir_client, base_url, "200000000004", qty=1, method="tunai")
    assert sale.get("cashier_id")


# ================================================================
# Excel export endpoints (owner only)
# ================================================================
def _assert_xlsx(resp):
    assert resp.status_code == 200, f"{resp.status_code} {resp.text[:200]}"
    assert resp.headers.get("content-type", "").startswith(XLSX_MIME)
    wb = load_workbook(io.BytesIO(resp.content))
    return wb


def test_export_sales_xlsx(owner_client, base_url):
    r = owner_client.get(f"{base_url}/api/export/sales?start=2026-01-01&end=2026-12-31")
    wb = _assert_xlsx(r)
    assert {"Ringkasan", "Transaksi", "Detail Item"}.issubset(set(wb.sheetnames))
    # Header row is row 4
    ws = wb["Transaksi"]
    headers = [c.value for c in ws[4]]
    assert "No. Transaksi" in headers and "Total" in headers
    # Row count matches JSON report
    rep = owner_client.get(f"{base_url}/api/reports/sales?start=2026-01-01&end=2026-12-31").json()
    data_rows = ws.max_row - 4
    assert data_rows == rep["summary"]["count"], f"rows {data_rows} != {rep['summary']['count']}"


def test_export_purchases_xlsx(owner_client, base_url):
    r = owner_client.get(f"{base_url}/api/export/purchases?start=2026-01-01&end=2026-12-31")
    wb = _assert_xlsx(r)
    assert {"Ringkasan", "Pembelian", "Detail Item"}.issubset(set(wb.sheetnames))
    ws = wb["Pembelian"]
    rep = owner_client.get(f"{base_url}/api/reports/purchases?start=2026-01-01&end=2026-12-31").json()
    assert ws.max_row - 4 == rep["summary"]["count"]


def test_export_stock_xlsx(owner_client, base_url):
    for mode in ("all", "low", "high"):
        r = owner_client.get(f"{base_url}/api/export/stock?mode={mode}")
        wb = _assert_xlsx(r)
        assert "Stok" in wb.sheetnames
        headers = [c.value for c in wb["Stok"][4]]
        assert "Kode" in headers and "Status" in headers


def test_export_mechanics_xlsx(owner_client, base_url):
    r = owner_client.get(f"{base_url}/api/export/mechanics?start=2026-01-01&end=2026-12-31")
    wb = _assert_xlsx(r)
    assert {"Ringkasan Mekanik", "Pekerjaan"}.issubset(set(wb.sheetnames))


def test_export_moves_xlsx(owner_client, base_url):
    r = owner_client.get(f"{base_url}/api/export/moves?start=2026-01-01&end=2026-12-31")
    wb = _assert_xlsx(r)
    assert "Mutasi Stok" in wb.sheetnames


def test_export_closings_xlsx(owner_client, base_url):
    r = owner_client.get(f"{base_url}/api/export/closings?start=2026-01-01&end=2026-12-31")
    wb = _assert_xlsx(r)
    assert "Tutup Kasir" in wb.sheetnames
    headers = [c.value for c in wb["Tutup Kasir"][4]]
    assert "No" in headers and "Kasir" in headers and "Selisih" in headers


def test_kasir_cannot_export(kasir_client, base_url):
    for path in ["/api/export/sales", "/api/export/purchases", "/api/export/stock",
                 "/api/export/mechanics", "/api/export/moves", "/api/export/closings"]:
        r = kasir_client.get(f"{base_url}{path}")
        assert r.status_code == 403, f"{path} expected 403, got {r.status_code}"
