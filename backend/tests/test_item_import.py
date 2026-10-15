# Tests for feature 19: Import barang dari Excel
import io
import os
import time
import requests
import pytest
from openpyxl import Workbook, load_workbook


def _build_xlsx(rows, header=None, header_at_row=1, extra_sheet=False):
    """rows: list of list; header defaults to template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Barang"
    hdr = header or ["Kode / Barcode", "Nama Barang *", "Kategori", "Satuan",
                     "Harga Beli", "Harga Jual *", "Stok Awal", "Stok Minimum"]
    # fill header at header_at_row (1-based)
    for i in range(1, header_at_row):
        ws.cell(row=i, column=1, value=f"filler {i}")
    for j, h in enumerate(hdr, start=1):
        ws.cell(row=header_at_row, column=j, value=h)
    for r_i, row in enumerate(rows, start=header_at_row + 1):
        for c_i, v in enumerate(row, start=1):
            ws.cell(row=r_i, column=c_i, value=v)
    if extra_sheet:
        wb.create_sheet("Petunjuk")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ---------- /api/items/import-template ----------
class TestImportTemplate:
    def test_template_download_owner(self, owner_token, base_url):
        r = requests.get(f"{base_url}/api/items/import-template",
                         headers={"Authorization": f"Bearer {owner_token}"}, timeout=30)
        assert r.status_code == 200
        assert "spreadsheet" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        assert "Barang" in wb.sheetnames
        assert "Petunjuk" in wb.sheetnames
        ws = wb["Barang"]
        # header should be at row 4
        headers_row = [c.value for c in ws[4]]
        expected = ["Kode / Barcode", "Nama Barang *", "Kategori", "Satuan",
                    "Harga Beli", "Harga Jual *", "Stok Awal", "Stok Minimum"]
        assert headers_row[:8] == expected
        # 2 example rows at 5 and 6
        row5 = [c.value for c in ws[5]]
        row6 = [c.value for c in ws[6]]
        assert row5[1] and row6[1]  # has name

    def test_template_forbidden_for_kasir(self, kasir_token, base_url):
        r = requests.get(f"{base_url}/api/items/import-template",
                         headers={"Authorization": f"Bearer {kasir_token}"}, timeout=30)
        assert r.status_code == 403


# ---------- /api/items/import ----------
class TestImportItems:
    def _post(self, base_url, token, content, dry_run=True, filename="test.xlsx"):
        files = {"file": (filename, content,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        return requests.post(
            f"{base_url}/api/items/import",
            params={"dry_run": str(dry_run).lower()},
            headers={"Authorization": f"Bearer {token}"},
            files=files, timeout=60,
        )

    def test_kasir_forbidden(self, kasir_token, base_url):
        content = _build_xlsx([["", "TEST_X", "", "", 1, 2, 0, 0]])
        r = self._post(base_url, kasir_token, content, dry_run=True)
        assert r.status_code == 403

    def test_invalid_file(self, owner_token, base_url):
        files = {"file": ("x.txt", b"not an xlsx", "text/plain")}
        r = requests.post(f"{base_url}/api/items/import",
                          params={"dry_run": "true"},
                          headers={"Authorization": f"Bearer {owner_token}"},
                          files=files, timeout=30)
        assert r.status_code == 400
        assert "tidak valid" in r.json().get("detail", "").lower()

    def test_file_too_large(self, owner_token, base_url):
        big = b"x" * (5_000_001)
        files = {"file": ("big.xlsx", big,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{base_url}/api/items/import",
                          params={"dry_run": "true"},
                          headers={"Authorization": f"Bearer {owner_token}"},
                          files=files, timeout=60)
        assert r.status_code == 400
        assert "5mb" in r.json().get("detail", "").lower()

    def test_dry_run_counts_and_rules(self, owner_token, base_url):
        suffix = str(int(time.time()))
        rows = [
            # empty code -> baru, with stock
            ["", f"TEST_Oli_{suffix}", "Oli", "botol", "45.000", "49.500", 5, 2],
            # existing code (demo item from seed) -> update, stock ignored
            ["200000000001", f"TEST_Update_{suffix}", "Oli", "botol", 1000, 2000, 99, 1],
            # missing name -> error
            ["", "", "x", "pcs", 1, 2, 0, 0],
            # missing sell_price -> error
            ["", f"TEST_NoPrice_{suffix}", "", "", 1, "", 0, 0],
            # non-numeric sell_price -> error
            ["", f"TEST_BadNum_{suffix}", "", "", 1, "abc", 0, 0],
            # negative -> error
            ["", f"TEST_Neg_{suffix}", "", "", -5, 10, 0, 0],
            # duplicate code within file -> second is error
            [f"TESTDUPX{suffix}", f"TEST_Dup1_{suffix}", "", "", 1, 2, 0, 0],
            [f"TESTDUPX{suffix}", f"TEST_Dup2_{suffix}", "", "", 1, 2, 0, 0],
        ]
        r = self._post(base_url, owner_token,
                       _build_xlsx(rows), dry_run=True)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["dry_run"] is True
        assert d["total"] == len(rows)
        # expected: baru = row1 + row7 = 2 ; update = row2 = 1 ; errors = 5
        assert d["created"] == 2
        assert d["updated"] == 1
        assert d["errors"] == 5

        by_row = {x["row"]: x for x in d["rows"]}
        # rows start at header+1 == row2 of excel; our header is row 1
        # so first data row = row 2
        actions = [by_row[2]["action"], by_row[3]["action"], by_row[4]["action"],
                   by_row[5]["action"], by_row[6]["action"], by_row[7]["action"],
                   by_row[8]["action"], by_row[9]["action"]]
        assert actions == ["baru", "update", "error", "error", "error", "error", "baru", "error"]
        # update message mentions stock not changed
        assert "stok" in by_row[3]["message"].lower()
        # dry run must NOT save: fetch item 200000000001 - name should NOT be TEST_Update_...
        got = requests.get(f"{base_url}/api/items",
                           headers={"Authorization": f"Bearer {owner_token}"},
                           params={"q": "200000000001"}, timeout=30)
        assert got.status_code == 200
        found = [i for i in got.json() if i["code"] == "200000000001"]
        assert found, "demo item should still exist"
        assert not found[0]["name"].startswith(f"TEST_Update_{suffix}")

    def test_header_detection_within_first_10_rows(self, owner_token, base_url):
        content = _build_xlsx(
            [["", f"TEST_HeaderRow_{int(time.time())}", "", "", 1, 2, 0, 0]],
            header_at_row=5,
        )
        r = self._post(base_url, owner_token, content, dry_run=True)
        assert r.status_code == 200
        d = r.json()
        assert d["total"] == 1
        assert d["created"] == 1

    def test_actual_import_creates_item_and_updates_existing(self, owner_token, base_url):
        suffix = str(int(time.time()))
        new_name = f"TEST_Import_New_{suffix}"
        upd_name = f"TEST_Import_Upd_{suffix}"
        rows = [
            ["", new_name, "Oli", "botol", "40.000", "49.500", 7, 2],  # baru + stock -> mutasi IMPORT
            ["200000000002", upd_name, "Rem", "pcs", 1234, 5678, 999, 3],  # update, stock ignored
        ]
        r = self._post(base_url, owner_token, _build_xlsx(rows), dry_run=False)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["created"] == 1 and d["updated"] == 1 and d["errors"] == 0

        # Verify new item exists
        got = requests.get(f"{base_url}/api/items",
                           headers={"Authorization": f"Bearer {owner_token}"},
                           params={"q": new_name}, timeout=30)
        assert got.status_code == 200
        new_items = [i for i in got.json() if i["name"] == new_name]
        assert new_items, "new item should be created"
        it = new_items[0]
        assert len(it["code"]) == 12 and it["code"].isdigit()
        assert it["sell_price"] == 49500  # Indonesian '49.500' parsed
        assert it["buy_price"] == 40000
        assert it["stock"] == 7
        assert it["min_stock"] == 2

        # Verify demo item 200000000002 was updated (name/price) but stock not changed to 999
        got2 = requests.get(f"{base_url}/api/items",
                            headers={"Authorization": f"Bearer {owner_token}"},
                            params={"q": "200000000002"}, timeout=30)
        upd = [i for i in got2.json() if i["code"] == "200000000002"][0]
        assert upd["name"] == upd_name
        assert upd["sell_price"] == 5678
        assert upd["stock"] != 999  # stock untouched

        # Verify stock movement IMPORT exists for the new item
        mv = requests.get(f"{base_url}/api/stock/moves",
                          headers={"Authorization": f"Bearer {owner_token}"},
                          params={"item_code": it["code"]}, timeout=30)
        assert mv.status_code == 200, mv.text
        moves = mv.json()
        imp = [m for m in moves if m.get("ref") == "IMPORT"]
        assert imp, f"expected an IMPORT mutasi for new item {it['code']}, got {moves[:3]}"
